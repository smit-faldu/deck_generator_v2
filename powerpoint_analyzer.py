import json
import os
from typing import Dict, List, Any, Optional
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor

class PowerPointAnalyzer:
    """
    A comprehensive tool for analyzing PowerPoint templates and extracting detailed information
    about their structure, content, and formatting.
    """
    
    def __init__(self, template_path: str):
        """
        Initialize the PowerPoint analyzer with a template path.
        
        Args:
            template_path: Path to the PowerPoint template file
        """
        self.template_path = template_path
        self.presentation = Presentation(template_path)
        self.analysis = {}
        
    def analyze(self) -> Dict:
        """
        Perform a comprehensive analysis of the PowerPoint template.
        
        Returns:
            Dict containing the analysis results
        """
        try:
            self.analysis = {
                'metadata': self._extract_metadata(),
                'slides': self._extract_slides(),
                'masters': self._extract_masters(),
                'themes': self._extract_themes(),
                'summary': self._generate_summary()
            }
            return self.analysis
        except Exception as e:
            print(f"Error during analysis: {str(e)}")
            # Return a minimal analysis with error information
            return {
                'error': str(e),
                'metadata': self._extract_metadata(),
                'slides': [],
                'masters': [],
                'themes': {'available': False, 'error': str(e)},
                'summary': {'error': str(e)}
            }
    
    def _extract_metadata(self) -> Dict:
        """Extract metadata about the presentation"""
        try:
            return {
                'slide_width': getattr(self.presentation, 'slide_width', 0),
                'slide_height': getattr(self.presentation, 'slide_height', 0),
                'core_properties': {
                    'author': getattr(self.presentation.core_properties, 'author', 'Unknown'),
                    'title': getattr(self.presentation.core_properties, 'title', 'Unknown'),
                    'subject': getattr(self.presentation.core_properties, 'subject', 'Unknown'),
                    'keywords': getattr(self.presentation.core_properties, 'keywords', 'Unknown'),
                    'created': str(getattr(self.presentation.core_properties, 'created', 'Unknown')),
                    'modified': str(getattr(self.presentation.core_properties, 'modified', 'Unknown')),
                }
            }
        except Exception as e:
            return {'error': str(e)}
    
    def _extract_slides(self) -> List[Dict]:
        """Extract detailed information about each slide"""
        slides = []
        
        try:
            for idx, slide in enumerate(self.presentation.slides):
                try:
                    slide_info = {
                        'index': idx,
                        'id': getattr(slide, 'slide_id', 'Unknown'),
                        'layout': getattr(slide.slide_layout, 'name', 'Unknown') if hasattr(slide, 'slide_layout') else 'Unknown',
                        'background': self._extract_background(slide),
                        'notes': self._extract_notes(slide),
                        'shapes': self._extract_shapes(slide),
                        'placeholders': self._extract_placeholders(slide)
                    }
                    slides.append(slide_info)
                except Exception as e:
                    print(f"Error extracting slide {idx}: {str(e)}")
                    slides.append({
                        'index': idx,
                        'error': str(e)
                    })
        except Exception as e:
            print(f"Error extracting slides: {str(e)}")
        
        return slides
    
    def _extract_masters(self) -> List[Dict]:
        """Extract information about slide masters"""
        masters = []
        
        try:
            for idx, master in enumerate(self.presentation.slide_masters):
                try:
                    master_info = {
                        'index': idx,
                        'name': getattr(master, 'name', f'Master {idx}'),
                        'shapes': self._extract_shapes(master),
                        'placeholders': self._extract_placeholders(master)
                    }
                    masters.append(master_info)
                except Exception as e:
                    print(f"Error extracting master {idx}: {str(e)}")
                    masters.append({
                        'index': idx,
                        'error': str(e)
                    })
        except Exception as e:
            print(f"Error extracting masters: {str(e)}")
        
        return masters
    
    def _extract_themes(self) -> Dict:
        """Extract theme information"""
        try:
            theme = getattr(self.presentation, 'theme', None)
            if not theme:
                return {'available': False}
            
            return {
                'available': True,
                'colors': self._extract_theme_colors(theme),
                'fonts': self._extract_theme_fonts(theme)
            }
        except Exception as e:
            print(f"Error extracting themes: {str(e)}")
            return {'available': False, 'error': str(e)}
    
    def _extract_theme_colors(self, theme) -> Dict:
        """Extract color information from the theme"""
        colors = {}
        
        try:
            if hasattr(theme, 'color_scheme') and hasattr(theme.color_scheme, '_element'):
                for color_scheme in theme.color_scheme._element:
                    if hasattr(color_scheme, 'name'):
                        name = color_scheme.name
                        if hasattr(color_scheme, 'val'):
                            colors[name] = color_scheme.val
        except Exception as e:
            print(f"Error extracting theme colors: {str(e)}")
        
        return colors
    
    def _extract_theme_fonts(self, theme) -> Dict:
        """Extract font information from the theme"""
        fonts = {}
        
        try:
            if hasattr(theme, 'font_scheme') and hasattr(theme.font_scheme, '_element'):
                for font_scheme in theme.font_scheme._element:
                    if hasattr(font_scheme, 'name'):
                        name = font_scheme.name
                        if hasattr(font_scheme, 'typeface'):
                            fonts[name] = font_scheme.typeface
        except Exception as e:
            print(f"Error extracting theme fonts: {str(e)}")
        
        return fonts
    
    def _extract_background(self, slide) -> Dict:
        """Extract background information from a slide"""
        try:
            background = slide.background
            if not background:
                return {'type': 'None'}
            
            fill = background.fill
            if not fill:
                return {'type': 'Empty'}
            
            return {
                'type': str(fill.type),
                'fore_color': self._extract_color(fill.fore_color),
                'back_color': self._extract_color(fill.back_color) if hasattr(fill, 'back_color') else {'type': 'None'}
            }
        except Exception as e:
            return {'type': 'Error', 'message': str(e)}
    
    def _extract_notes(self, slide) -> Dict:
        """Extract notes information from a slide"""
        try:
            if not hasattr(slide, 'notes_slide') or not slide.notes_slide:
                return {'available': False}
            
            notes_slide = slide.notes_slide
            notes_text = notes_slide.notes_text_frame.text if hasattr(notes_slide, 'notes_text_frame') else ''
            
            return {
                'available': True,
                'text': notes_text,
                'shapes': self._extract_shapes(notes_slide)
            }
        except Exception as e:
            return {'available': False, 'error': str(e)}
    
    def _extract_shapes(self, slide) -> List[Dict]:
        """Extract detailed information about shapes on a slide"""
        shapes = []
        
        try:
            for shape in slide.shapes:
                try:
                    shape_info = {
                        'id': getattr(shape, 'shape_id', 'Unknown'),
                        'name': getattr(shape, 'name', 'Unknown'),
                        'type': str(getattr(shape, 'shape_type', 'Unknown')),
                        'position': {
                            'left': getattr(shape, 'left', 0),
                            'top': getattr(shape, 'top', 0),
                            'width': getattr(shape, 'width', 0),
                            'height': getattr(shape, 'height', 0),
                            'rotation': getattr(shape, 'rotation', 0)
                        },
                        'text': self._extract_text(shape),
                        'fill': self._extract_fill(shape),
                        'line': self._extract_line(shape),
                        'special_content': self._extract_special_content(shape)
                    }
                    shapes.append(shape_info)
                except Exception as e:
                    print(f"Error extracting shape: {str(e)}")
                    shapes.append({
                        'error': str(e)
                    })
        except Exception as e:
            print(f"Error extracting shapes: {str(e)}")
        
        return shapes
    
    def _extract_text(self, shape) -> Dict:
        """Extract text and formatting information from a shape"""
        try:
            if not hasattr(shape, 'text_frame') or not shape.has_text_frame:
                return {'available': False}
            
            text_frame = shape.text_frame
            paragraphs = []
            
            for para in text_frame.paragraphs:
                try:
                    para_info = {
                        'text': para.text,
                        'alignment': str(getattr(para, 'alignment', 'None')),
                        'level': getattr(para, 'level', 0),
                        'runs': []
                    }
                    
                    for run in para.runs:
                        try:
                            run_info = {
                                'text': run.text,
                                'font': {
                                    'name': getattr(run.font, 'name', 'None'),
                                    'size': getattr(run.font.size, 'pt', 'None') if hasattr(run.font, 'size') and run.font.size else 'None',
                                    'bold': getattr(run.font, 'bold', 'None'),
                                    'italic': getattr(run.font, 'italic', 'None'),
                                    'underline': getattr(run.font, 'underline', 'None'),
                                    'color': self._extract_color(getattr(run.font, 'color', None))
                                }
                            }
                            para_info['runs'].append(run_info)
                        except Exception as e:
                            print(f"Error extracting run: {str(e)}")
                            para_info['runs'].append({'error': str(e)})
                    
                    paragraphs.append(para_info)
                except Exception as e:
                    print(f"Error extracting paragraph: {str(e)}")
                    paragraphs.append({'error': str(e)})
            
            return {
                'available': True,
                'text': shape.text,
                'paragraphs': paragraphs,
                'word_wrap': getattr(text_frame, 'word_wrap', 'None'),
                'auto_size': getattr(text_frame, 'auto_size', 'None')
            }
        except Exception as e:
            return {'available': False, 'error': str(e)}
    
    def _extract_fill(self, shape) -> Dict:
        """Extract fill information from a shape"""
        try:
            if not hasattr(shape, 'fill'):
                return {'available': False}
            
            fill = shape.fill
            if not fill:
                return {'available': False, 'type': 'None'}
            
            return {
                'available': True,
                'type': str(fill.type),
                'fore_color': self._extract_color(fill.fore_color),
                'back_color': self._extract_color(fill.back_color) if hasattr(fill, 'back_color') else {'type': 'None'},
                'transparency': fill.transparency if hasattr(fill, 'transparency') else 'None'
            }
        except Exception as e:
            return {'available': False, 'type': 'Error', 'message': str(e)}
    
    def _extract_line(self, shape) -> Dict:
        """Extract line information from a shape"""
        try:
            if not hasattr(shape, 'line'):
                return {'available': False}
            
            line = shape.line
            if not line:
                return {'available': False, 'type': 'None'}
            
            return {
                'available': True,
                'type': str(line.type),
                'color': self._extract_color(line.color),
                'width': line.width if hasattr(line, 'width') else 'None',
                'dash_style': str(line.dash_style) if hasattr(line, 'dash_style') else 'None'
            }
        except Exception as e:
            return {'available': False, 'type': 'Error', 'message': str(e)}
    
    def _extract_color(self, color) -> Dict:
        """Extract color information"""
        try:
            if not color:
                return {'type': 'None'}
            
            if hasattr(color, 'rgb'):
                return {
                    'type': 'RGB',
                    'value': str(color.rgb)
                }
            elif hasattr(color, 'theme_color'):
                return {
                    'type': 'Theme',
                    'value': str(color.theme_color)
                }
            else:
                return {'type': 'Unknown'}
        except Exception as e:
            return {'type': 'Error', 'message': str(e)}
    
    def _extract_special_content(self, shape) -> Dict:
        """Extract information about special content types (tables, charts, etc.)"""
        try:
            # Check for tables
            if hasattr(shape, 'has_table') and shape.has_table:
                table = shape.table
                return {
                    'type': 'Table',
                    'rows': len(table.rows),
                    'columns': len(table.columns),
                    'cells': [[cell.text for cell in row.cells] for row in table.rows]
                }
            
            # Check for charts
            if hasattr(shape, 'chart'):
                chart = shape.chart
                return {
                    'type': 'Chart',
                    'chart_type': str(chart.chart_type),
                    'title': chart.chart_title.text_frame.text if hasattr(chart, 'chart_title') else 'None',
                    'series_count': len(chart.series) if hasattr(chart, 'series') else 0
                }
            
            # Check for pictures
            if hasattr(shape, 'image'):
                return {
                    'type': 'Picture',
                    'format': shape.image.ext if hasattr(shape.image, 'ext') else 'Unknown',
                    'size': {
                        'width': shape.width,
                        'height': shape.height
                    }
                }
            
            # Check for group shapes
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                return {
                    'type': 'Group',
                    'shape_count': len(shape.shapes) if hasattr(shape, 'shapes') else 0
                }
            
            return {'type': 'None'}
        except Exception as e:
            return {'type': 'Error', 'message': str(e)}
    
    def _extract_placeholders(self, slide) -> List[Dict]:
        """Extract information about placeholders on a slide"""
        placeholders = []
        
        try:
            for shape in slide.placeholders:
                try:
                    placeholder_info = {
                        'id': getattr(shape.placeholder_format, 'idx', 'Unknown'),
                        'type': str(getattr(shape.placeholder_format, 'type', 'Unknown')),
                        'name': getattr(shape, 'name', 'Unknown'),
                        'text': getattr(shape, 'text', ''),
                        'position': {
                            'left': getattr(shape, 'left', 0),
                            'top': getattr(shape, 'top', 0),
                            'width': getattr(shape, 'width', 0),
                            'height': getattr(shape, 'height', 0)
                        }
                    }
                    placeholders.append(placeholder_info)
                except Exception as e:
                    print(f"Error extracting placeholder: {str(e)}")
                    placeholders.append({'error': str(e)})
        except Exception as e:
            print(f"Error extracting placeholders: {str(e)}")
        
        return placeholders
    
    def _generate_summary(self) -> Dict:
        """Generate a summary of the presentation"""
        try:
            slides = self.analysis.get('slides', [])
            
            return {
                'slide_count': len(slides),
                'master_count': len(self.analysis.get('masters', [])),
                'placeholder_count': sum(len(slide.get('placeholders', [])) for slide in slides),
                'shape_count': sum(len(slide.get('shapes', [])) for slide in slides),
                'text_count': sum(1 for slide in slides for shape in slide.get('shapes', []) 
                                 if shape.get('text', {}).get('available', False)),
                'table_count': sum(1 for slide in slides for shape in slide.get('shapes', []) 
                                  if shape.get('special_content', {}).get('type') == 'Table'),
                'chart_count': sum(1 for slide in slides for shape in slide.get('shapes', []) 
                                  if shape.get('special_content', {}).get('type') == 'Chart'),
                'picture_count': sum(1 for slide in slides for shape in slide.get('shapes', []) 
                                    if shape.get('special_content', {}).get('type') == 'Picture')
            }
        except Exception as e:
            print(f"Error generating summary: {str(e)}")
            return {'error': str(e)}
    
    def save_analysis(self, output_path: str) -> None:
        """
        Save the analysis results to a JSON file.
        
        Args:
            output_path: Path to save the JSON file
        """
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(self.analysis, f, indent=2, default=str)
    
    def print_summary(self) -> None:
        """Print a summary of the analysis results"""
        summary = self.analysis.get('summary', {})
        
        print(f"\n=== PowerPoint Template Analysis Summary ===")
        print(f"Template: {os.path.basename(self.template_path)}")
        print(f"Slides: {summary.get('slide_count', 0)}")
        print(f"Masters: {summary.get('master_count', 0)}")
        print(f"Placeholders: {summary.get('placeholder_count', 0)}")
        print(f"Shapes: {summary.get('shape_count', 0)}")
        print(f"Text elements: {summary.get('text_count', 0)}")
        print(f"Tables: {summary.get('table_count', 0)}")
        print(f"Charts: {summary.get('chart_count', 0)}")
        print(f"Pictures: {summary.get('picture_count', 0)}")
        print(f"==========================================\n")


def main():
    """Example usage of the PowerPointAnalyzer"""
    import argparse
    
    parser = argparse.ArgumentParser(description='Analyze a PowerPoint template')
    parser.add_argument('--template', '-t', default='template/Black Elegant and Modern Startup Pitch Deck Presentation.pptx',
                        help='Path to the PowerPoint template (default: template/Black Elegant and Modern Startup Pitch Deck Presentation.pptx)')
    parser.add_argument('--output', '-o', default='template_analysis.json',
                        help='Path to save the analysis results (JSON) (default: template_analysis.json)')
    parser.add_argument('--summary', '-s', action='store_true', help='Print a summary of the analysis')
    
    args = parser.parse_args()
    
    # Check if template file exists
    if not os.path.exists(args.template):
        print(f"Error: Template file not found at '{args.template}'")
        print("Please provide a valid template path using --template or -t")
        return
    
    try:
        print(f"Analyzing template: {args.template}")
        analyzer = PowerPointAnalyzer(args.template)
        analysis = analyzer.analyze()
        
        if args.summary:
            analyzer.print_summary()
        
        analyzer.save_analysis(args.output)
        print(f"Analysis saved to {args.output}")
    except Exception as e:
        print(f"Error analyzing template: {str(e)}")


if __name__ == "__main__":
    main() 